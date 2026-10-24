"""Read the document being checked into blocks: one per text frame,
paragraph, table cell or chart value, each with a location a reader can find.
Pictures are collected alongside, for the model to transcribe (images.py)."""

from __future__ import annotations

import re
from collections import Counter
from collections.abc import Iterable
from dataclasses import dataclass, field
from pathlib import Path

from onus import charts

# A block that is only a short number is a page number or a list marker.
_BARE_NUMBER = re.compile(r"^\s*\d{1,3}\s*$")


@dataclass(frozen=True)
class Block:
    id: int
    label: str  # "slide 3", "paragraph 12", "line 40"
    text: str
    plain: bool = False  # ends in a chart value printed with no format: a magnitude, no unit


@dataclass(frozen=True)
class Figure:
    label: str  # the slide or paragraph it sits on, as a Block's label
    ordinal: int
    data: bytes
    media_type: str

    @property
    def location(self) -> str:
        return f"{self.label} (image {self.ordinal})"


@dataclass
class Document:
    path: Path
    format: str
    blocks: list[Block]
    figures: list[Figure] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n".join(b.text for b in self.blocks)

    def extend(self, items: Iterable[tuple[str, str, bool]]) -> None:
        seen = {(b.label, b.text) for b in self.blocks}
        for label, text, plain in items:
            text = " ".join(text.split())
            if not text or _BARE_NUMBER.match(text):
                continue
            if (label, text) in seen:  # a slide repeats its eyebrow, a table its header
                continue
            seen.add((label, text))
            self.blocks.append(Block(len(self.blocks), label, text, plain))


def _shapes(shapes):
    from pptx.shapes.group import GroupShape

    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _shapes(shape.shapes)
        else:
            yield shape


def _pptx(path: Path) -> tuple[list[tuple[str, str, bool]], list[Figure]]:
    from pptx import Presentation
    from pptx.shapes.picture import Picture

    texts: list[tuple[str, str, bool]] = []
    figures: list[Figure] = []
    for n, slide in enumerate(Presentation(str(path)).slides, start=1):
        pictures = graphs = 0
        for shape in _shapes(slide.shapes):
            if getattr(shape, "has_chart", False) and shape.has_chart:
                graphs += 1
                texts += charts.blocks(shape.chart, f"slide {n} (chart {graphs})")
                continue
            if isinstance(shape, Picture):
                pictures += 1
                rId = shape._element.blip_rId
                if rId is None:  # a linked picture has no bytes in the file
                    continue
                part = shape.part.related_part(rId)
                figures.append(Figure(f"slide {n}", pictures, part.blob, part.content_type))
                continue
            frames = []
            if shape.has_text_frame:
                frames = [shape.text_frame]
            elif getattr(shape, "has_table", False) and shape.has_table:
                frames = [c.text_frame for row in shape.table.rows for c in row.cells]
            for frame in frames:
                for paragraph in frame.text.split("\n"):
                    texts.append((f"slide {n}", paragraph, False))
    return texts, figures


def _docx(path: Path) -> tuple[list[tuple[str, str, bool]], list[Figure]]:
    import docx

    d = docx.Document(str(path))
    texts: list[tuple[str, str, bool]] = []
    figures: list[Figure] = []
    pictures: Counter[str] = Counter()

    def collect(paragraphs, label: str) -> None:
        for p in paragraphs:
            for drawing in p._p.xpath(".//w:drawing"):
                rids = drawing.xpath(".//a:blip/@r:embed")
                if not rids:  # a native chart or SmartArt, not a picture
                    continue
                pictures[label] += 1
                part = d.part.related_parts[rids[0]]
                figures.append(Figure(label, pictures[label], part.blob, part.content_type))

    for i, p in enumerate(d.paragraphs, start=1):
        texts.append((f"paragraph {i}", p.text, False))
        collect([p], f"paragraph {i}")
    for t, table in enumerate(d.tables, start=1):
        seen = []
        for row in table.rows:
            for cell in row.cells:
                texts.append((f"table {t}", cell.text, False))
                if cell._tc not in seen:  # a merged cell is listed once per column it spans
                    seen.append(cell._tc)
                    collect(cell.paragraphs, f"table {t}")
    return texts, figures


def _text(path: Path) -> tuple[list[tuple[str, str, bool]], list[Figure]]:
    return [(f"line {i}", line, False) for i, line in enumerate(path.read_text().splitlines(), start=1)], []


READERS = {".pptx": _pptx, ".docx": _docx, ".md": _text, ".txt": _text, ".markdown": _text}


def load(path: str | Path) -> Document:
    path = Path(path)
    reader = READERS.get(path.suffix.lower())
    if reader is None:
        raise ValueError(f"{path.name}: unsupported format (onus reads {', '.join(sorted(READERS))})")
    texts, figures = reader(path)
    doc = Document(path, path.suffix.lower().lstrip("."), [], figures)
    doc.extend(texts)
    return doc
